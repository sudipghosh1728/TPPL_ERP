"""Bounded Excel import and QC report exports. Workbook content is data, never code."""
import base64,datetime as dt,io,re,zipfile
from decimal import Decimal,InvalidOperation

STEPS=('Cutting','Machining','Tapping','Finishing','QC','Blasting','Primer & Engraving','Final Paint','Final QC','Packing')

def text(value):
 if value is None:return ''
 if isinstance(value,float):return format(value,'.12g')
 if isinstance(value,(dt.date,dt.datetime)):return value.date().isoformat() if isinstance(value,dt.datetime) else value.isoformat()
 return str(value).strip()

def file_bytes(file):
 if not isinstance(file,dict) or not str(file.get('name','')).lower().endswith('.xlsx'):raise ValueError('Choose an .xlsx workbook')
 try:raw=base64.b64decode(file.get('content',''),validate=True)
 except Exception:raise ValueError('Invalid Excel file')
 if not 0<len(raw)<=3*1024*1024:raise ValueError('Workbook must be below 3 MB')
 try:
  with zipfile.ZipFile(io.BytesIO(raw)) as z:
   if len(z.infolist())>300 or sum(x.file_size for x in z.infolist())>20*1024*1024:raise ValueError('Workbook expands beyond the import limit')
   for item in z.infolist():
    if item.filename.endswith('.xml') and (b'<!DOCTYPE' in z.read(item) or b'<!ENTITY' in z.read(item)):raise ValueError('XML entities are not allowed')
 except zipfile.BadZipFile:raise ValueError('Invalid Excel workbook')
 return raw

def workbook(file):
 from openpyxl import load_workbook
 wb=load_workbook(io.BytesIO(file_bytes(file)),read_only=False,data_only=False,keep_links=False)
 if len(wb.worksheets)>12 or any(s.max_row>2000 or s.max_column>40 for s in wb):wb.close();raise ValueError('Workbook exceeds supported sheet dimensions')
 return wb

def date(value):
 if isinstance(value,(dt.date,dt.datetime)):return text(value)[:10]
 s=text(value)
 for pattern in ('%Y-%m-%d','%d.%m.%Y','%d/%m/%Y'):
  try:return dt.datetime.strptime(s,pattern).date().isoformat()
  except ValueError:pass
 if isinstance(value,(int,float)) and 30000<=value<=80000:
  from openpyxl.utils.datetime import from_excel
  return from_excel(value).date().isoformat()
 return ''

def integer(value,maximum=100000):
 try:
  n=Decimal(str(value))
  if not n.is_finite() or n!=int(n) or not 0<n<=maximum:raise ValueError()
  return int(n)
 except (ValueError,TypeError,InvalidOperation,OverflowError):raise ValueError('Enter a positive whole quantity (maximum '+str(maximum)+')')

def tracker(file):
 wb=workbook(file)
 try:
  if 'PROCESS' not in wb.sheetnames:raise ValueError('Tracker must contain the PROCESS sheet')
  ws=wb['PROCESS'];headers=[text(ws.cell(1,c).value).casefold() for c in range(1,19)]
  if headers[:3]!=['po','drawing no','thickness'] or headers[13]!='qty':raise ValueError('PROCESS columns do not match the PO / Drawing No / Thickness / stages / Qty format')
  batches=[];warnings=[]
  for row in range(2,ws.max_row+1):
   vals=[ws.cell(row,c).value for c in range(1,19)]
   if not vals[0] or not vals[1]:continue
   drg=text(vals[1]);match=re.match(r'([\w.-]+)',drg)
   if not match:raise ValueError('Invalid drawing at PROCESS row '+str(row))
   planned=date(vals[14]);actual=date(vals[15])
   if not planned:warnings.append('PROCESS row '+str(row)+': planned dispatch date needs confirmation')
   batches.append(dict(source_row=row,po=text(vals[0]),drawing=match[1],description=drg,thickness=text(vals[2]),quantity=integer(vals[13],100),planned_date=planned,source_actual_date=actual,drawing_ref=text(vals[16]),numbering_ref=text(vals[17]),source_steps={step:text(vals[i+3]) for i,step in enumerate(STEPS)}))
  if not batches or len(batches)>200:raise ValueError('Import between 1 and 200 process rows at a time')
  trace=[]
  if 'Current PO' in wb.sheetnames:
   for row in range(2,wb['Current PO'].max_row+1):
    vals=[wb['Current PO'].cell(row,c).value for c in range(1,8)]
    if not vals[4]:continue
    trace.append(dict(source_row=row,thickness=text(vals[0]),heat_no=text(vals[1]),plate_no=text(vals[2]),ai_no=text(vals[3]),quantity=integer(vals[4]),planned_date=date(vals[5]),source_actual_date=date(vals[6])))
  warnings+=['Imported progress and dates are source claims; QC must verify them before release.','Current PO has no drawing column: allocate heat/plate quantities explicitly; no drawing association is guessed.','Drawing and numbering references are filenames, not embedded drawings. Upload the actual files against each batch.','OLD PO and other sheets are retained in the original workbook; they are not silently merged with current production.']
  return dict(kind='tracker',batches=batches,trace_lots=trace,warnings=warnings)
 finally:wb.close()

NUM=r'-?\d+(?:\.\d+)?'
def limits(spec,tolerance):
 if not text(tolerance):
  embedded=re.fullmatch(r'(.*?)\(\s*((?:±|\+/-)\s*\d+(?:\.\d+)?)\s*\)',text(spec))
  if embedded:spec,tolerance=embedded.groups()
  else:
   bound=re.fullmatch(r'('+NUM+r')\s*(MAX|MIN)',text(spec),re.I)
   if bound:return [None,str(Decimal(bound[1]))] if bound[2].upper()=='MAX' else [str(Decimal(bound[1])),None]
 nominal=re.fullmatch(r'(?:[ØøRr]\s*)?('+NUM+r')(?:\s*\([^)]*\))?',text(spec))
 tol=re.fullmatch(r'\+/-\s*(\d+(?:\.\d+)?)',text(tolerance).lstrip('.').replace('±','+/-'))
 if not nominal or not tol:return None
 n=Decimal(nominal[1]);t=Decimal(tol[1]);return [str(n-t),str(n+t)]

def cell_result(characteristic,value):
 s=text(value)
 if not s:return 'Missing'
 bounds=characteristic.get('limits')
 if bounds:
  if not re.fullmatch(NUM,s):return 'Review'
  n=Decimal(s);return 'OK' if (bounds[0] is None or Decimal(bounds[0])<=n) and (bounds[1] is None or n<=Decimal(bounds[1])) else 'Rework'
 return characteristic.get('manual_result') if characteristic.get('manual_result') in ('OK','Rework') and characteristic.get('manual_reason') else 'Review'

def results(report):
 counts={k:0 for k in ('OK','Rework','Review','Missing')}
 for row in report['characteristics']:
  for value in row['values']:counts[cell_result(row,value)]+=1
 return counts

def inspection_quantity(ws):
 """Read a labelled quantity without truncating decimals or executing formulas."""
 candidates=[]
 for cells in ws.iter_rows(min_row=1,max_row=5,max_col=min(ws.max_column,40)):
  for cell in cells:
   match=re.fullmatch(r'(?:QTY|QUANTITY)\b[\s.:#=-]*(.*)',text(cell.value),re.I)
   if not match:continue
   value=match[1].strip()
   if not value:
    # Quantity may be in a separate cell to the right of its label.
    for col in range(cell.column+1,min(cell.column+3,ws.max_column)+1):
     adjacent=text(ws.cell(cell.row,col).value)
     if adjacent:
      if re.fullmatch(r'\d+(?:\.\d+)?(?:\s*(?:NOS?\.?|PCS?\.?|PIECES?))?',adjacent,re.I):value=adjacent
      break
   if value:candidates.append((cell.coordinate,value))
 if not candidates and isinstance(ws['J4'].value,(int,float)) and not isinstance(ws['J4'].value,bool):candidates=[('J4',text(ws['J4'].value))]
 quantities=[]
 for coordinate,value in candidates:
  cleaned=re.sub(r'\s*(?:NOS?\.?|PCS?\.?|PIECES?)$','',value,flags=re.I).strip()
  try:qty=integer(cleaned,100)
  except ValueError:raise ValueError(f'{ws.title}!{coordinate}: quantity must be a whole number from 1 to 100 (for example QTY: 20).')
  quantities.append(qty)
 if len(set(quantities))>1:raise ValueError(ws.title+': conflicting quantity labels; confirm the inspection quantity.')
 return quantities[0] if quantities else None

def inspection(file,blank=False,quantity=None,po=None):
 wb=workbook(file)
 try:
  rows={};samples=set();header=None;warnings=[]
  for ws in wb:
   if 'INSPECTION' not in text(ws['C1'].value).upper():continue
   def label(cell,pattern):
    matches=[]
    for row in ws.iter_rows(min_row=3,max_row=4,max_col=min(ws.max_column,40)):
     for value in row:
      match=re.match(pattern,text(value.value),re.I)
      if match and match[1].strip():matches.append(match[1].strip())
    if len(set(matches))>1:raise ValueError(ws.title+': conflicting header values near '+cell)
    return matches[0] if matches else ''
   try:source_quantity=inspection_quantity(ws)
   except ValueError as error:
    if not blank or quantity is None:raise
    source_quantity=None;warnings.append(str(error)+' Blank import uses the selected batch quantity instead.')
   target_quantity=integer(quantity,100) if blank and quantity is not None else source_quantity
   if target_quantity is None:raise ValueError(ws.title+': inspection quantity is missing. Enter QTY: 20 in the workbook header, or use Blank format from a drawing batch to use its quantity.')
   source_po=label('D4',r'PO\s*No[.\s:-]*(\w[\w.-]*)')
   if blank and po is not None and not source_po:warnings.append(ws.title+': blank report uses the selected project PO because the workbook PO is empty.')
   if blank and quantity is not None and source_quantity!=target_quantity:warnings.append(ws.title+': blank report uses the selected batch quantity ('+str(target_quantity)+'); workbook quantity is '+(str(source_quantity) if source_quantity else 'blank')+'.')
   meta=dict(customer=label('A3',r'CUSTOMER\s*[:-]+\s*(.+)'),part_name=label('D3',r'Part Name[.\s:-]*(.+)'),drawing=label('A4',r'DRG\s*No\s*[:-]+\s*([\w.-]+)'),revision=label('C4',r'REV\s*[:-]+\s*(.+)'),po=str(po) if blank and po is not None else source_po,quantity=target_quantity,inspection_date=date(label('J3',r'Date\s*[:-]+\s*(.+)')))
   if not meta['drawing'] or not meta['po']:raise ValueError('Drawing and PO are required in every inspection sheet')
   if header and any(meta[k]!=header[k] for k in ('drawing','po','quantity','revision')):raise ValueError('Inspection sheets have conflicting drawing, PO, revision or quantity')
   header=header or meta
   def piece(value):
    if value is None or text(value)=='':return None
    try:return integer(value,100)
    except ValueError:return None
   first_col,last_col,remark_col=5,9,10
   for cells in ws.iter_rows(min_row=5,max_row=7,max_col=min(ws.max_column,40)):
    actual=next((v for v in cells if re.fullmatch(r'ACTUAL(?:\s+(?:SIZE|READINGS?|VALUES?))?',text(v.value),re.I)),None)
    remark=next((v for v in cells if re.fullmatch(r'REMARKS?',text(v.value),re.I)),None)
    if actual and remark and remark.column>actual.column:
     first_col,last_col,remark_col=actual.column,remark.column-1,remark.column;break
   sample_cols=range(first_col,last_col+1)
   header_row=next((r for r in (7,8) if ws.cell(r,1).value in (None,'') and any(piece(ws.cell(r,c).value) for c in sample_cols) and all(not text(ws.cell(r,c).value) or piece(ws.cell(r,c).value) for c in sample_cols)),None)
   if header_row is None and blank:
    header_row=next((r-1 for r in range(7,13) if piece(ws.cell(r,1).value) and text(ws.cell(r,3).value)),None)
   if header_row is None:raise ValueError('Cannot locate piece numbers in '+ws.title)
   columns=[]
   numbered=[piece(ws.cell(header_row,c).value) for c in sample_cols if piece(ws.cell(header_row,c).value)]
   if blank and source_quantity and numbered and max(numbered)>source_quantity:warnings.append(ws.title+': piece numbering exceeds the workbook quantity. Blank import discards the source readings; confirm the batch quantity before inspection.')
   for col in (() if blank else sample_cols):
    if not text(ws.cell(header_row,col).value):
     if any(piece(ws.cell(r,1).value) and text(ws.cell(r,col).value) for r in range(header_row+1,min(ws.max_row,150)+1)):raise ValueError(ws.title+': measurements have no piece number in column '+ws.cell(header_row,col).column_letter+'. Add the piece number before historical import.')
     continue
    sample=integer(ws.cell(header_row,col).value,100)
    if sample in samples or sample>meta['quantity']:raise ValueError(ws.title+': piece '+str(sample)+' is duplicated or exceeds header quantity '+str(meta['quantity'])+'. Correct the workbook before importing historical readings.')
    samples.add(sample);columns.append((col,sample))
   reference='';section='';sheet_rows=set()
   for r in range(header_row+1,min(ws.max_row,150)+1):
    sr=ws.cell(r,1).value
    try:key=str(integer(sr,10000))
    except ValueError:
     context=text(ws.cell(r,2).value)
     if sr is None and context and text(ws.cell(r,3).value) in ('',context):
      if context.lower().startswith('reference'):reference=context;section=''
      else:section=context
     continue
    if key in sheet_rows:raise ValueError(ws.title+': duplicate characteristic '+key)
    sheet_rows.add(key);spec=text(ws.cell(r,3).value);tolerance=text(ws.cell(r,4).value) if first_col>4 else ''
    description=' · '.join(v for v in (reference,section,text(ws.cell(r,2).value)) if v)
    row=rows.setdefault(key,dict(number=key,description=description,specified=spec,tolerance=tolerance,limits=limits(spec,tolerance),values=['']*meta['quantity'],source_remarks=[],manual_result='',manual_reason=''))
    if row['specified']!=spec or row['tolerance']!=tolerance:raise ValueError('Specification differs between sheets at characteristic '+key)
    for col,sample in columns:row['values'][sample-1]='' if blank else text(ws.cell(r,col).value)
    remark=text(ws.cell(r,remark_col).value)
    if remark and remark not in row['source_remarks']:row['source_remarks'].append(remark)
  if not header or not rows or len(rows)>120:raise ValueError('No supported inspection matrix found')
  filename=str(file.get('name','')).rsplit('.',1)[0]
  if re.fullmatch(r'\d+',filename) and filename!=header['drawing']:warnings.append('Filename '+filename+' differs from drawing '+header['drawing']+' in the workbook header. Verify the drawing reference; the header is used for matching.')
  if not blank and samples!=set(range(1,header['quantity']+1)):warnings.append('Some piece columns are absent; missing measurements remain empty.')
  report={**header,'characteristics':[rows[k] for k in sorted(rows,key=int)],'warnings':warnings}
  report['summary']=results(report)
  if report['summary']['Rework']:warnings.append(str(report['summary']['Rework'])+' measurements exceed the stated tolerance, regardless of the original OK remark.')
  if report['summary']['Review']:warnings.append('Text readings / unrecognised tolerances need QC review; no automatic acceptance was inferred.')
  return dict(kind='inspection',report=report,warnings=warnings)
 finally:wb.close()

def export_xlsx(report,batch,trace,company):
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill,Alignment
 from openpyxl.utils import get_column_letter
 wb=Workbook();wb.remove(wb.active)
 for start in range(0,report['quantity'],5):
  ws=wb.create_sheet('Pieces '+str(start+1)+'-'+str(min(start+5,report['quantity'])))
  ws.append([company+' · INSPECTION REPORT']);ws.append(['Report',report['reference'],'Status',report['status'],'Revision',report['revision']])
  ws.append(['Customer',report['customer'],'Part',report['part_name']]);ws.append(['PO',report['po'],'Drawing',report['drawing'],'Date',report['inspection_date']])
  ws.append(['Inspector',report.get('inspector',''),'Approved by',report.get('approved_by',''),'Approved at',report.get('approved_at','')])
  ws.append(['No.','Description','Specified','Tolerance',*[str(i+1) for i in range(start,min(start+5,report['quantity']))],'Result'])
  for row in report['characteristics']:
   values=row['values'][start:start+5];verdicts=[cell_result(row,v) for v in values]
   ws.append([row['number'],row['description'],row['specified'],row['tolerance'],*values,', '.join(dict.fromkeys(verdicts))])
   for i,v in enumerate(verdicts,5):ws.cell(ws.max_row,i).fill=PatternFill('solid',fgColor='DCFCE7' if v=='OK' else 'FEE2E2' if v=='Rework' else 'FEF3C7')
  ws.append(['Review notes',report.get('review_notes','')]);ws.append(['Approval notes',report.get('approval_notes','')]);ws.append(['Source remarks are retained separately; this report uses the current evaluated measurements.'])
  ws.freeze_panes='E7';ws.print_title_rows='1:6';ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
 ws=wb.create_sheet('Traceability');ws.append(['Heat No','Plate No','AI No','Quantity','Thickness'])
 for t in trace:ws.append([t.get('heat_no',''),t.get('plate_no',''),t.get('ai_no',''),t['allocated'],t.get('thickness','')])
 ws=wb.create_sheet('Process');ws.append(['Step','Status','Completed quantity','Notes'])
 for step in STEPS:
  s=batch['steps'][step];ws.append([step,s['status'],s['quantity'],s.get('notes','')])
 ws.append(['Process evidence at report approval' if report['status']=='Approved' else 'Current process evidence for this draft'])
 ws=wb.create_sheet('Assessments');ws.append(['Characteristic','Manual result','Reason / method','Assessed by','Original Excel remarks'])
 for row in report['characteristics']:ws.append([row['number'],row.get('manual_result',''),row.get('manual_reason',''),row.get('manual_by',''),', '.join(row.get('source_remarks',[]))])
 for ws in wb:
  for cells in ws:
   for cell in cells:
    if isinstance(cell.value,str):cell.data_type='s' # Never turn user text into spreadsheet formulas.
    cell.alignment=Alignment(vertical='top',wrap_text=True)
  for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17664F')
  for col in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(col)].width=22 if col<5 else 16
 output=io.BytesIO();wb.save(output);return output.getvalue()

def export_pdf(report,batch,trace,company):
 from html import escape
 from reportlab.lib import colors
 from reportlab.lib.pagesizes import A4,landscape
 from reportlab.lib.styles import getSampleStyleSheet
 from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,LongTable,TableStyle,PageBreak
 output=io.BytesIO();styles=getSampleStyleSheet();styles['BodyText'].fontSize=8;styles['BodyText'].leading=11
 def p(s):return Paragraph(escape(text(s)),styles['BodyText'])
 story=[]
 for start in range(0,report['quantity'],5):
  if start:story.append(PageBreak())
  story+=[Paragraph(escape(company)+' · Inspection report',styles['Title']),p(report['reference']+' · '+report['status']+' · Drawing '+report['drawing']+' Rev '+report['revision']),p('PO '+report['po']+' · '+report['customer']+' · '+report['part_name']),p('Date '+report['inspection_date']+' · Inspector '+report.get('inspector','')+' · Approved by '+report.get('approved_by','Not approved')),Spacer(1,10)]
  data=[list(map(p,['No','Description','Specified','Tolerance',*['Piece '+str(i+1) for i in range(start,min(start+5,report['quantity']))],'Result']))]
  for row in report['characteristics']:
   vals=row['values'][start:start+5];data.append(list(map(p,[row['number'],row['description'],row['specified'],row['tolerance'],*vals,', '.join(dict.fromkeys(cell_result(row,v) for v in vals))])))
  widths=[30,65,80,65]+[65]*len(range(start,min(start+5,report['quantity'])))+[80]
  table=LongTable(data,colWidths=widths,repeatRows=1);table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dceee5')),('GRID',(0,0),(-1,-1),.3,colors.lightgrey),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),5)]));story.append(table)
  story+=[Spacer(1,8),p('Review notes: '+report.get('review_notes','')),p('Approval notes: '+report.get('approval_notes',''))]
 story+=[PageBreak(),Paragraph('Material traceability',styles['Heading1'])]
 for row in trace:story.append(p('Heat '+row.get('heat_no','')+' · Plate '+row.get('plate_no','')+' · AI '+row.get('ai_no','')+' · Qty '+str(row['allocated'])))
 story+=[Spacer(1,12),Paragraph('Process at report approval' if report['status']=='Approved' else 'Current process record',styles['Heading1'])]
 for step in STEPS:
  s=batch['steps'][step];story.append(p(step+' · '+s['status']+' · Qty '+str(s['quantity'])+' · '+s.get('notes','')))
 story+=[Spacer(1,12),Paragraph('Manual assessments',styles['Heading1'])]
 for row in report['characteristics']:
  if row.get('manual_result'):story.append(p('Characteristic '+row['number']+' · '+row['manual_result']+' · '+row.get('manual_reason','')+' · '+row.get('manual_by','')))
 SimpleDocTemplate(output,pagesize=landscape(A4),rightMargin=28,leftMargin=28,topMargin=28,bottomMargin=28).build(story)
 return output.getvalue()
