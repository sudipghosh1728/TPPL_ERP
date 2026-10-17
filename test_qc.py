import base64,copy,io,json,unittest,zipfile
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
import access,seed,server,qc_excel as x,qc_ops as q,project_flow as pf
from mongo_test_support import test_store,clean

def workbook_file(header=8):
 wb=Workbook();wb.remove(wb.active)
 for start in (1,6):
  ws=wb.create_sheet(str(start));ws['C1']='INSPECTION REPORT';ws['A3']='CUSTOMER:- Test customer';ws['D3']='Part Name.- Test plate';ws['J3']='Date:- 01.09.2026'
  ws['A4']='DRG No:- TEST-01';ws['C4']='REV - 00';ws['D4']='PO No. :- PO-QC';ws['J4']='QTY:- 10'
  for col,n in enumerate(range(start,start+5),5):ws.cell(header,col,n)
  for r,values in enumerate([(1,'A','Ø32','.+/-0.2','32.2'),(2,'B','Visual','As drawing','No damage')],header+1):
   for col,value in enumerate(values[:4],1):ws.cell(r,col,value)
   for col in range(5,10):ws.cell(r,col,values[4])
   ws.cell(r,10,'OK')
 out=io.BytesIO();wb.save(out);return dict(name='inspection.xlsx',content=base64.b64encode(out.getvalue()).decode())

def tracker_file():
 wb=Workbook();ws=wb.active;ws.title='PROCESS';ws.append(['PO','Drawing No','Thickness',*x.STEPS,'Qty','Dispatch Date','Actual Dispatch Date','Drawing','Numbering Ref'])
 for day in ('2026-09-01','2026-10-01'):ws.append(['PO-QC','TEST-01 (Machining)','20mm',*(['Done']*10),10,day,day,'test-drawing.pdf','test-numbering.jpeg'])
 ws=wb.create_sheet('Current PO');ws.append(['Thickness','Heat No','Plate No','AI number','Qty','Dispatch Date','Actual Dispatch Date']);ws.append(['20mm','HEAT-1','PLATE-1','AI-1',12,'2026-09-01','']);ws.append(['20mm','HEAT-2','PLATE-2','AI-2',8,'2026-10-01',''])
 out=io.BytesIO();wb.save(out);return dict(name='tracker.xlsx',content=base64.b64encode(out.getvalue()).decode())

class QCExcelTests(unittest.TestCase):
 def test_layouts_limits_formulas_and_blank_templates(self):
  for header in (7,8):
   r=x.inspection(workbook_file(header))['report'];self.assertEqual(r['quantity'],10);self.assertEqual(len(r['characteristics']),2);self.assertEqual(r['summary'],{'OK':10,'Review':10,'Rework':0,'Missing':0})
   self.assertEqual(x.cell_result(r['characteristics'][0],'32.21'),'Rework');self.assertEqual(x.cell_result(r['characteristics'][0],'31.8'),'OK');self.assertEqual(x.cell_result(r['characteristics'][0],'32 .1'),'Review');self.assertEqual(x.cell_result(r['characteristics'][0],'=32+0.1'),'Review')
   self.assertEqual(x.inspection(workbook_file(header),blank=True)['report']['summary']['Missing'],20)
  t=x.tracker(tracker_file());self.assertEqual(len(t['batches']),2);self.assertEqual(len(t['trace_lots']),2);self.assertNotIn('drawing',t['trace_lots'][0])
  with self.assertRaises(ValueError):x.file_bytes({'name':'bad.xlsx','content':base64.b64encode(b'not a zip').decode()})

class QCWorkflowTests(unittest.TestCase):
 def setUp(self):
  self.store=test_store();self.store.prepare();self.store.run(lambda c:seed.initialize(c,demo=True));self.old=server.STORE;server.STORE=self.store
  self.users={role:dict(id=i+1,name=role+' employee',role=role,username=role.lower(),active=1) for i,role in enumerate(access.ROLES)}
  self.store.run(lambda c:pf.create(c,dict(plant_id='solu',title='QC project',customer='Test customer',po_number='PO-QC',quantity=20,due='2026-12-01',route='Fabrication only',file=dict(name='po.txt',content='UE8=')),self.users['Admin']))
  self.jid=self.store.run(lambda c:c.all('work_jobs',sort=[('id',-1)])[0]['id'])
 def tearDown(self):server.STORE=self.old;clean(self.store)
 def rows(self,name):return self.store.run(lambda c:c.all(name,{'job_id':self.jid},sort=[('id',1)]))
 def call(self,role,path,**d):
  if d.get('batch_id') and 'version' not in d:d['version']=self.store.run(lambda c:c.one('qc_batches',{'id':d['batch_id']}))['version']
  with patch('server.access.current',return_value=self.users[role]):return self.store.run(lambda c:server.write_response(c,'/api/qc/'+path,d,None,'test'))
 def test_every_user_can_prepare_reports_but_only_qc_admin_approve(self):
  for role in access.ROLES:
   self.call(role,'create',job_id=self.jid,drawing='TEST-01',thickness='20mm',quantity=10,planned_date='2026-12-01')
   b=self.rows('qc_batches')[-1];self.call(role,'import-report',batch_id=b['id'],file=workbook_file(),mode='historical');r=self.rows('qc_reports')[-1]
   for char in r['characteristics']:
    if not char['limits']:char.update(manual_result='OK',manual_reason='Visual check against drawing')
   self.call(role,'save-report',batch_id=b['id'],report_version=r['version'],inspection_date='2026-09-01',characteristics=r['characteristics'],review_notes='Prepared by '+role)
   if role not in ('QC','Admin'):
    with self.assertRaises(access.AccessError):self.call(role,'approve-report',batch_id=b['id'],report_version=2,notes='Not authorised')
   pdf=self.store.run(lambda c:q.export(c,r['id'],self.users[role],'pdf'))['_download']['content'];self.assertTrue(pdf.startswith(b'%PDF'))
  self.assertEqual(len(self.rows('qc_reports')),len(access.ROLES))
 def test_stale_edits_evidence_changes_and_blank_reuse(self):
  self.call('HR','create',job_id=self.jid,drawing='TEST-01',thickness='20mm',quantity=10,planned_date='2026-12-01')
  b=self.rows('qc_batches')[0];self.call('HR','import-report',batch_id=b['id'],file=workbook_file(),mode='historical')
  r=self.rows('qc_reports')[0];chars=copy.deepcopy(r['characteristics']);chars[1].update(manual_result='OK',manual_reason='Visual inspection')
  chars[0]['limits']=['0','1000']
  with self.assertRaises(ValueError):self.call('HR','save-report',batch_id=b['id'],version=b['version'],report_version=1,inspection_date='2026-09-01',characteristics=chars)
  self.call('HR','save-report',batch_id=b['id'],report_version=1,inspection_date='2026-09-01',characteristics=chars)
  saved=self.rows('qc_reports')[0];self.assertEqual(saved['characteristics'][0]['limits'],['31.8','32.2'])
  with self.assertRaises(ValueError):self.call('HR','save-report',batch_id=b['id'],report_version=1,inspection_date='2026-09-01',characteristics=chars)
  with self.assertRaises(access.AccessError):self.call('HR','trace',batch_id=b['id'],new_lot=dict(heat_no='H',plate_no='P',quantity=10))
  self.call('QC','trace',batch_id=b['id'],new_lot=dict(heat_no='H',plate_no='P',quantity=10));lot=self.rows('qc_trace_lots')[0]
  self.call('QC','trace',batch_id=b['id'],allocations=[dict(lot_id=lot['id'],quantity=10)])
  self.call('QC','document',batch_id=b['id'],kind='Drawing',file=dict(name='drawing.pdf',content='VGVzdA=='))
  self.call('QC','approve-report',batch_id=b['id'],report_version=2,notes='Approved')
  self.call('QC','document',batch_id=b['id'],kind='TC',file=dict(name='tc.pdf',content='VGVzdA=='))
  errors=self.store.run(lambda c:q.release_errors(c,c.one('qc_batches',{'id':b['id']})))
  self.assertTrue(any('Evidence changed' in e for e in errors))
  self.call('Accounts','revise-report',batch_id=b['id'],notes='Certificate revision')
  reports=self.rows('qc_reports');self.assertEqual(reports[0]['status'],'Approved');self.assertEqual(reports[1]['status'],'Draft');self.assertEqual(reports[1]['revises'],r['id'])
  self.call('Maintenance','import-report',batch_id=b['id'],template_report_id=r['id'])
  blank=self.rows('qc_reports')[-1];self.assertTrue(all(not any(v for v in row['values']) and not row['manual_result'] for row in blank['characteristics']))
 def test_trace_limits_release_snapshots_and_dispatch(self):
  source=tracker_file()
  self.call('QC','import-tracker',job_id=self.jid,file=source,request_id='same-import')
  self.call('QC','import-tracker',job_id=self.jid,file=source,request_id='same-import');self.assertEqual(len(self.rows('qc_batches')),2)
  with self.assertRaises(ValueError):self.call('QC','import-tracker',job_id=self.jid,file=source)
  b,other=self.rows('qc_batches');lots=self.rows('qc_trace_lots');self.assertTrue(all(s['status']=='Pending' for s in b['steps'].values()));self.assertEqual(b['actual_date'],'')
  with self.assertRaises(ValueError):self.store.run(lambda c:q.gate(c,c.one('work_jobs',{'id':self.jid})))
  self.call('QC','trace',batch_id=b['id'],allocations=[{'lot_id':lots[0]['id'],'quantity':10}])
  with self.assertRaises(ValueError):self.call('QC','trace',batch_id=other['id'],allocations=[{'lot_id':lots[0]['id'],'quantity':10}])
  self.call('Store','import-report',batch_id=b['id'],file=workbook_file(),mode='historical');r=self.rows('qc_reports')[0]
  with self.assertRaises(ValueError):self.call('QC','approve-report',batch_id=b['id'],report_version=1,notes='Missing manual review')
  chars=copy.deepcopy(r['characteristics']);chars[1].update(manual_result='OK',manual_reason='Drawing visual criteria verified');chars[0]['values'][0]='32.21'
  self.call('Store','save-report',batch_id=b['id'],report_version=1,inspection_date='2026-09-01',characteristics=chars)
  with self.assertRaises(ValueError):self.call('QC','approve-report',batch_id=b['id'],report_version=2,notes='Failing dimension')
  chars[0]['values'][0]='32.19';self.call('Store','save-report',batch_id=b['id'],report_version=2,inspection_date='2026-09-01',characteristics=chars)
  for kind in ('Drawing','TC','Numbering'):self.call('QC','document',batch_id=b['id'],kind=kind,file=dict(name=kind+'.pdf',content='VGVzdA=='))
  self.call('QC','approve-report',batch_id=b['id'],report_version=3,notes='Verified')
  before=self.store.run(lambda c:q.export(c,r['id'],self.users['Store'],'xlsx'))['_download']['content']
  workbook=load_workbook(io.BytesIO(before));self.assertEqual(workbook.worksheets[0]['F2'].value,'00');self.assertIn('Traceability',workbook.sheetnames)
  with self.assertRaises(ValueError):self.call('Store','save-report',batch_id=b['id'],report_version=4,inspection_date='2026-09-01',characteristics=chars)
  steps={s:dict(status='Done',quantity=10,notes='Verified') for s in x.STEPS}
  with self.assertRaises(ValueError):self.call('QC','steps',batch_id=b['id'],steps=steps,notes='Packing cannot precede final inspection')
  steps['Packing']=dict(status='Pending',quantity=0,notes='Awaiting final inspection')
  self.call('QC','steps',batch_id=b['id'],steps=steps,notes='Ready for project final inspection')
  self.assertEqual(self.store.run(lambda c:q.release_errors(c,c.one('qc_batches',{'id':b['id']}),include_packing=False)),[])
  self.store.run(lambda c:c.update('work_jobs',{'id':self.jid},{'stage':'Packing & dispatch','final_inspection':True}))
  steps['Packing']=dict(status='Done',quantity=10,notes='Verified')
  self.call('QC','steps',batch_id=b['id'],steps=steps,notes='Process and packing complete')
  self.assertEqual(self.store.run(lambda c:q.release_errors(c,c.one('qc_batches',{'id':b['id']}))),[])
  self.call('QC','document',batch_id=b['id'],kind='EWAY challan',file=dict(name='eway.pdf',content='VGVzdA=='))
  after=self.store.run(lambda c:q.export(c,r['id'],self.users['Store'],'xlsx'))['_download']['content']
  self.assertEqual(load_workbook(io.BytesIO(before))['Process']['B11'].value,load_workbook(io.BytesIO(after))['Process']['B11'].value)
  self.store.run(lambda c:c.update('work_jobs',{'id':self.jid},{'stage':'QC inspection','final_inspection':False}))
  with self.assertRaises(ValueError):self.call('QC','dispatch',batch_id=b['id'],actual_date='2026-09-02',notes='Project gate still closed')
  self.store.run(lambda c:c.update('work_jobs',{'id':self.jid},{'stage':'Packing & dispatch','final_inspection':True}))
  self.call('QC','dispatch',batch_id=b['id'],actual_date='2026-09-02',notes='Loaded')
  packet=self.store.run(lambda c:q.packet(c,b['id'],self.users['HR']))['_download']['content'];names=zipfile.ZipFile(io.BytesIO(packet)).namelist();self.assertTrue(any(n.endswith('.pdf') for n in names));self.assertIn('traceability.json',names)
  with self.assertRaises(ValueError):self.call('HR','revise-report',batch_id=b['id'],notes='Dispatched record')

if __name__=='__main__':unittest.main()
