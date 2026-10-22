import base64,copy,io,unittest,zipfile
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill
import access,seed,server,report_sheet as sheet,report_workshop as reports,project_flow as pf,project_deletion,qc_ops as qc
from mongo_test_support import test_store,clean

def fixture():
 w=Workbook();s=w.active;s.title='Inspection';s.merge_cells('A1:D1');s['A1']='CUSTOM REPORT';s['A1'].font=Font(bold=True,size=16);s['A1'].fill=PatternFill('solid',fgColor='CCDDFF');s['A3']=2;s['A4']=3;s['A5']='=SUM(A3:A4)';s['B5']="='Other sheet'!A1*2";w.create_sheet('Other sheet')['A1']=7
 out=io.BytesIO();w.save(out)
 return out.getvalue()

class SheetTests(unittest.TestCase):
 def test_layout_style_formula_and_sheet_roundtrip(self):
  raw=fixture();m=sheet.read(raw);self.assertEqual(m['sheets'][0]['merges'],['A1:D1']);self.assertEqual(sheet.calculate(m)['s1'],{'A5':5.0,'B5':14.0})
  m['sheets'][0]['cells']['A1']['fmt']={'bold':False,'fill':'#123456','align':'center'}
  m['sheets'][0]['cells']['A3']['fmt']={'number_format':'0.00%'}
  m['sheets'].append(dict(id='new',source='',name='Paint',rows=20,cols=10,cells={'A1':dict(v='Painting remarks',t='s',s=0)},merges=[],widths={},heights={},orientation='landscape'))
  exported=sheet.export(raw,sheet.validate(m,sheet.read(raw)));w=load_workbook(io.BytesIO(exported))
  self.assertEqual(w.sheetnames,['Inspection','Other sheet','Paint']);self.assertEqual(w['Paint']['A1'].value,'Painting remarks');self.assertFalse(w['Inspection']['A1'].font.bold);self.assertEqual(w['Inspection']['A1'].fill.fgColor.rgb,'FF123456');w.close()
  self.assertEqual(w['Inspection']['A3'].number_format,'0.00%')
  w=load_workbook(io.BytesIO(exported),data_only=True);self.assertEqual(w['Inspection']['A5'].value,5);w.close()
 def test_invalid_merges_source_formulas_and_xml(self):
  raw=fixture();original=sheet.read(raw)
  for update in (lambda m:m['sheets'][0].update(source='../../secret'),lambda m:m['sheets'][0]['merges'].append('B1:E2'),lambda m:m['sheets'][0]['cells']['A1'].update(v='bad\x00')):
   m=copy.deepcopy(original);update(m)
   with self.assertRaises(ValueError):sheet.validate(m,original)
  m=copy.deepcopy(original);m['sheets'][0]['cells']['A5']['v']='=A5';m['sheets'][0]['cells']['B5']['v']='=__import__("os").system("unsafe")';result=sheet.calculate(m)
  self.assertEqual(result['s1']['A5'],'#CYCLE!');self.assertEqual(result['s1']['B5'],'#EXCEL!')
 def test_embedded_original_parts_are_retained(self):
  raw=fixture();data=sheet.parts(raw);data['xl/media/vector.emf']=b'vector-evidence';data['xl/drawings/evidence.xml']=b'<drawing />';out=io.BytesIO()
  with zipfile.ZipFile(out,'w') as z:
   for name,value in data.items():z.writestr(name,value)
  raw=out.getvalue();exported=sheet.parts(sheet.export(raw,sheet.read(raw)))
  self.assertEqual(exported['xl/media/vector.emf'],b'vector-evidence');self.assertEqual(exported['xl/drawings/evidence.xml'],b'<drawing />')

class WorkshopWorkflowTests(unittest.TestCase):
 def setUp(self):
  self.store=test_store();self.store.prepare();self.store.run(lambda c:seed.initialize(c,demo=True));self.old=server.STORE;server.STORE=self.store
  self.users={role:dict(id=i+1,name=role+' employee',role=role,username=role.lower(),active=1) for i,role in enumerate(access.ROLES)}
 def tearDown(self):server.STORE=self.old;clean(self.store)
 def call(self,role,path,**data):
  with patch('server.access.current',return_value=self.users[role]):return self.store.run(lambda c:server.write_response(c,'/api/workbooks/'+path,data,None,'test'))
 def row(self,rid):return self.store.run(lambda c:c.one('report_workbooks',{'id':rid}))
 def create(self,role='Store',**extra):return self.call(role,'create',name='Custom report',file={'name':'format.xlsx','content':base64.b64encode(fixture()).decode()},**extra)['id']
 def test_all_departments_save_review_lock_revise_and_conflicts(self):
  for role in access.ROLES:
   rid=self.create(role);row=self.row(rid);row['model']['sheets'][0]['cells']['A3']['v']='8'
   self.call(role,'save',id=rid,version=1,model=row['model'])
   history=self.store.run(lambda c:c.one('report_workbook_audit',{'report_id':rid,'action':'Saved draft'}))
   self.assertEqual(history['previous_model']['sheets'][0]['cells']['A3']['v'],'2')
   with self.assertRaisesRegex(ValueError,'Another user'):self.call(role,'save',id=rid,version=1,model=row['model'])
   if role not in ('Admin','QC'):
    with self.assertRaises(access.AccessError):self.call(role,'approve',id=rid,version=2,notes='Unauthorized')
   self.call('QC','approve',id=rid,version=2,notes='Reviewed all sheets')
   with self.assertRaisesRegex(ValueError,'locked'):self.call(role,'save',id=rid,version=3,model=row['model'])
   revision=self.call(role,'duplicate',id=rid,version=3)['id'];self.assertEqual(self.row(revision)['status'],'Draft');self.assertEqual(self.row(rid)['status'],'Approved')
 def test_project_permissions_rename_and_archive(self):
  self.store.run(lambda c:pf.create(c,dict(plant_id='solu',title='Report project',customer='Test',po_number='PO-R',quantity=1,due='2026-12-01',route='Fabrication only',file=dict(name='po.txt',content='UE8=')),self.users['Admin']))
  jid=self.store.run(lambda c:c.all('work_jobs',sort=[('id',-1)])[0]['id']);rid=self.create(job_id=jid);row=self.row(rid)
  row['model']['sheets'][1]['name']='Renamed';self.call('Store','save',id=rid,version=1,model=row['model']);self.assertEqual(sheet.calculate(self.row(rid)['model'])['s1']['B5'],14)
  self.store.run(lambda c:c.update('work_jobs',{'id':jid},{'assignments':{'QC':999}}))
  with self.assertRaises(access.AccessError):self.call('QC','approve',id=rid,version=2,notes='Not assigned')
  self.call('Admin','approve',id=rid,version=2,notes='Admin review')
  data,queries,digest=self.store.run(lambda c:project_deletion.snapshot(c,jid,self.users['Admin']))
  packet=self.store.run(lambda c:project_deletion.make_zip(c,data,digest,self.users['Admin']))
  with zipfile.ZipFile(io.BytesIO(packet)) as z:self.assertEqual(z.read(f'report-workshop/{rid}-original.xlsx'),self.row(rid)['source']);self.assertIn('report_workbooks',queries)
  self.store.run(lambda c:c.update('work_jobs',{'id':jid},{'stage':'Completed'}))
  with self.assertRaisesRegex(ValueError,'read-only'):self.call('Store','duplicate',id=rid,version=3)

 def test_upload_any_layout_preserves_bytes_roles_revisions_and_archive(self):
  for role in access.ROLES:
   raw=b'%PDF-1.4\nCompleted external report for '+role.encode()
   rid=self.call(role,'upload',name='Paint report',category='Painting',file={'name':'paint.pdf','content':base64.b64encode(raw).decode()})['id']
   row=self.row(rid);self.assertEqual(row['mode'],'uploaded');self.assertNotIn('model',row)
   self.assertEqual(self.store.run(lambda c:reports.export(c,rid,self.users[role]))['_download']['content'],raw)
   if role not in ('Admin','QC'):
    with self.assertRaises(access.AccessError):self.call(role,'approve',id=rid,version=1,notes='Invalid reviewer')
   self.call('QC','approve',id=rid,version=1,notes='Downloaded and reviewed')
   revised=self.call(role,'upload',id=rid,version=2,name='Paint report revision',category='Painting',file={'name':'revised.xlsx','content':base64.b64encode(b'Any worksheet layout; no parsing').decode()})['id']
   self.assertEqual(self.row(revised)['status'],'Draft');self.assertEqual(self.row(revised)['revision_of'],rid);self.assertEqual(self.row(rid)['source'],raw);self.assertEqual(self.row(rid)['status'],'Approved')
   with self.assertRaises(ValueError):self.call(role,'save',id=revised,version=1,model={})
  for name,raw in [('unsafe.exe',b'bad'),('empty.pdf',b''),('large.pdf',b'x'*(3*1024*1024+1))]:
   with self.assertRaises(ValueError):self.call('Store','upload',name='Invalid',file={'name':name,'content':base64.b64encode(raw).decode()})

 def test_batch_upload_review_evidence_and_zip(self):
  admin=self.users['Admin'];reviewer=self.users['QC'];author=self.users['Store']
  self.store.run(lambda c:pf.create(c,dict(plant_id='solu',title='Upload project',customer='Test',po_number='PO-U',quantity=1,due='2026-12-01',route='Fabrication only',file=dict(name='po.txt',content='UE8=')),admin))
  jid=self.store.run(lambda c:c.all('work_jobs',sort=[('id',-1)])[0]['id'])
  self.store.run(lambda c:qc.create_batch(c,dict(job_id=jid,drawing='ANY-01',thickness='20mm',quantity=1,planned_date='2026-12-01'),author))
  def batch():return self.store.run(lambda c:c.all('qc_batches',{'job_id':jid})[0])
  file={'name':'final-report.pdf','content':base64.b64encode(b'%PDF-1.4 arbitrary finished report').decode()}
  b=batch();rid=self.call('Store','upload',name='Inspection',category='Inspection',job_id=jid,batch_id=b['id'],batch_version=b['version'],file=file)['id']
  b=batch();self.assertEqual(b['current_upload'],rid)
  with self.assertRaisesRegex(ValueError,'traceability'):self.call('QC','approve',id=rid,version=1,batch_version=b['version'],notes='Incomplete evidence')
  self.store.run(lambda c:qc.trace(c,dict(batch_id=b['id'],version=b['version'],new_lot=dict(heat_no='H1',plate_no='P1',quantity=1)),reviewer));b=batch()
  lot=self.store.run(lambda c:c.all('qc_trace_lots',{'job_id':jid})[0])
  self.store.run(lambda c:qc.trace(c,dict(batch_id=b['id'],version=b['version'],allocations=[dict(lot_id=lot['id'],quantity=1)]),reviewer));b=batch()
  self.store.run(lambda c:qc.document(c,dict(batch_id=b['id'],version=b['version'],kind='Drawing',file={'name':'drawing.pdf','content':'UERG'}),reviewer));b=batch()
  self.call('QC','approve',id=rid,version=1,batch_version=b['version'],notes='Checked uploaded report against drawing and heat trace')
  b=batch();self.assertEqual(self.store.run(lambda c:qc.release_errors(c,b,check_steps=False)),[])
  packet=self.store.run(lambda c:qc.packet(c,b['id'],author))['_download']['content']
  with zipfile.ZipFile(io.BytesIO(packet)) as z:self.assertEqual(z.read(f'uploaded-reports/{rid}-Approved-final-report.pdf'),base64.b64decode(file['content']))
  old_version=b['version'];self.store.run(lambda c:qc.schedule(c,dict(batch_id=b['id'],version=b['version'],planned_date='2026-12-02',notes='Changed date'),reviewer))
  with self.assertRaisesRegex(ValueError,'changed'):self.call('Store','upload',id=rid,version=2,name='Revision',category='Inspection',batch_version=old_version,file=file)
  b=batch();revision=self.call('Store','upload',id=rid,version=2,name='Revision',category='Inspection',batch_version=b['version'],file=file)['id'];b=batch()
  self.assertEqual(b['current_upload'],revision);self.assertIn('Approve the current inspection report',self.store.run(lambda c:qc.release_errors(c,b,check_steps=False)))
  self.assertEqual(self.row(rid)['status'],'Approved')
  self.store.run(lambda c:c.update('work_jobs',{'id':jid},{'stage':'Completed'}))
  with self.assertRaisesRegex(ValueError,'read-only'):self.call('Store','upload',id=revision,version=1,name='Invalid revision',file=file)

if __name__=='__main__':unittest.main()
