"""Layout-independent report workbooks. Patch OOXML to retain original drawings/media."""
import ast,base64,copy,datetime as dt,io,json,math,posixpath,re,warnings,zipfile
import xml.etree.ElementTree as ET
from openpyxl import Workbook,load_workbook
from openpyxl.utils import get_column_letter,column_index_from_string,range_boundaries
from qc_excel import file_bytes

S='http://schemas.openxmlformats.org/spreadsheetml/2006/main';R='http://schemas.openxmlformats.org/officeDocument/2006/relationships';P='http://schemas.openxmlformats.org/package/2006/relationships';C='http://schemas.openxmlformats.org/package/2006/content-types'
ET.register_namespace('',S);ET.register_namespace('r',R)
def tag(name):return '{'+S+'}'+name
def xml(raw):
 if b'<!DOCTYPE' in raw or b'<!ENTITY' in raw:raise ValueError('XML entities are not allowed')
 return ET.fromstring(raw)
def encoded(node):
 # ElementTree assigns new prefixes; stale mc:Ignorable prefix lists cause Excel repairs.
 for item in node.iter():item.attrib.pop('{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable',None)
 return ET.tostring(node,encoding='utf-8',xml_declaration=True)
def target(base,path):return path.lstrip('/') if path.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(base),path))
def parts(raw):
 with zipfile.ZipFile(io.BytesIO(raw)) as z:return {n:z.read(n) for n in z.namelist()}
def blank():
 w=Workbook();w.active.title='Report';out=io.BytesIO();w.save(out);return out.getvalue()
def rgb(col,default='000000'):
 return '#'+col.rgb[-6:] if col and col.type=='rgb' and isinstance(col.rgb,str) else '#'+default
def position(ref):
 if not isinstance(ref,str) or not re.fullmatch(r'[A-Z]{1,3}[1-9]\d{0,3}',ref):raise ValueError('Invalid cell address')
 match=re.fullmatch(r'([A-Z]+)(\d+)',ref);col=column_index_from_string(match[1]);row=int(match[2])
 if row>1000 or col>100:raise ValueError('Use up to 1000 rows and 100 columns per sheet')
 return row,col
def relationships(data,path):
 rel=posixpath.join(posixpath.dirname(path),'_rels',posixpath.basename(path)+'.rels')
 return {r.get('Id'):target(path,r.get('Target')) for r in xml(data[rel]) if r.get('TargetMode')!='External'} if rel in data else {}
def pictures(data,sheet_path):
 out=[];notes=[];rels=relationships(data,sheet_path);root=xml(data[sheet_path])
 for drawing in root.findall(tag('drawing')):
  path=rels.get(drawing.get('{'+R+'}id'));drawing_rels=relationships(data,path) if path else {}
  if not path or path not in data:continue
  for anchor in xml(data[path]):
   blip=anchor.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
   if blip is None:continue
   image_path=drawing_rels.get(blip.get('{'+R+'}embed'));raw=data.get(image_path,b'')
   mime='image/png' if raw.startswith(b'\x89PNG') else 'image/jpeg' if raw.startswith(b'\xff\xd8') else None
   if not mime:
    notes.append('An EMF/WMF or unsupported drawing is retained in the Excel download; browser/PDF previews cannot display it.');continue
   def number(group,key,default=0):
    element=anchor.find('{*}'+group+'/{*}'+key)
    return int(element.text) if element is not None else default
   ext=anchor.find('{*}ext');col=number('from','col');row=number('from','row')
   out.append(dict(src='data:'+mime+';base64,'+base64.b64encode(raw).decode(),col=col,row=row,end_col=number('to','col',col+2),end_row=number('to','row',row+4),width=int(ext.get('cx'))/9525 if ext is not None else None,height=int(ext.get('cy'))/9525 if ext is not None else None,dx=number('from','colOff')/9525,dy=number('from','rowOff')/9525))
 return out,list(dict.fromkeys(notes))

def sizes(ws):
 widths={str(i):min(600,max(35,round((ws.sheet_format.defaultColWidth or 13)*7))) for i in range(1,ws.max_column+1)}
 for dim in list(ws.column_dimensions.values()):
  start=dim.min or column_index_from_string(dim.index);end=dim.max or start
  for i in range(start,min(end,ws.max_column)+1):widths[str(i)]=min(600,max(35,round((dim.width or 13)*7)))
 heights={str(i):min(500,max(20,round(((ws.row_dimensions.get(i).height if ws.row_dimensions.get(i) else None) or ws.sheet_format.defaultRowHeight or 15)*4/3))) for i in range(1,ws.max_row+1)}
 return widths,heights

def read(raw):
 data=parts(raw);root=xml(data['xl/workbook.xml']);rels=relationships(data,'xl/workbook.xml');paths=[rels[s.get('{'+R+'}id')] for s in root.find(tag('sheets'))]
 with warnings.catch_warnings():
  warnings.simplefilter('ignore',UserWarning);w=load_workbook(io.BytesIO(raw),data_only=False,keep_links=False)
 if len(w.worksheets)>50:raise ValueError('Use up to 50 sheets per report workbook')
 styles={};sheets=[]
 try:
  for index,ws in enumerate(w):
   if ws.max_row>1000 or ws.max_column>100:raise ValueError('Workbook exceeds 1000 rows / 100 columns per sheet')
   cells={}
   for row in ws:
    for c in row:
     if c.value is None and not c.has_style:continue
     sid=str(c.style_id)
     if sid not in styles:styles[sid]=dict(bold=bool(c.font.bold),italic=bool(c.font.italic),size=c.font.sz or 11,color=rgb(c.font.color),fill=rgb(c.fill.fgColor,'FFFFFF') if c.fill.patternType=='solid' else '#ffffff',align=c.alignment.horizontal or 'left',wrap=bool(c.alignment.wrap_text),border=any(side and side.style for side in (c.border.left,c.border.right,c.border.top,c.border.bottom)))
     styles[sid]['number_format']=c.number_format
     value=c.value;typ=c.data_type
     if isinstance(value,(dt.datetime,dt.date)):value=value.isoformat();typ='d'
     cells[c.coordinate]=dict(v='' if value is None else str(value),t='s' if value is None else typ,s=c.style_id)
   cols,rows=sizes(ws)
   sheets.append(dict(id='s'+str(index+1),source=paths[index],name=ws.title,rows=max(ws.max_row,20),cols=max(ws.max_column,8),cells=cells,merges=[str(r) for r in ws.merged_cells.ranges],widths=cols,heights=rows,orientation=ws.page_setup.orientation or 'portrait'))
  return dict(sheets=sheets,styles=styles)
 finally:w.close()

def validate(model,original):
 if not isinstance(model,dict) or len(json.dumps(model,ensure_ascii=False).encode())>2500000:raise ValueError('Report workbook exceeds the editor size limit (2.5 MB)')
 sheets=model.get('sheets');allowed={s['source'] for s in original['sheets']};style_ids=set(original['styles'])|{'0'}
 if not isinstance(sheets,list) or not 1<=len(sheets)<=50:raise ValueError('Keep between 1 and 50 sheets')
 # Original sheets stay in the file, including their drawings and private evidence.
 if {s.get('source') for s in sheets if s.get('source')}!=allowed:raise ValueError('Original sheets must be retained; clear cells or rename sheets instead')
 names=set();ids=set();sources=set();total=0;out=[]
 for sheet in sheets:
  name=str(sheet.get('name',''));sid=str(sheet.get('id',''));source=sheet.get('source','')
  if not name.strip() or len(name)>31 or re.search(r'[\\/*?:\[\]]',name) or name.casefold() in names:raise ValueError('Sheet names must be unique and valid Excel names (31 characters maximum)')
  if not re.fullmatch(r'[A-Za-z0-9_-]{1,50}',sid) or sid in ids or (source and (source not in allowed or source in sources)):raise ValueError('Invalid sheet identity')
  names.add(name.casefold());ids.add(sid);sources.add(source)
  nr=sheet.get('rows');nc=sheet.get('cols')
  if type(nr)!=int or type(nc)!=int or not 1<=nr<=1000 or not 1<=nc<=100:raise ValueError('Use 1–1000 rows and 1–100 columns')
  cells={}
  for ref,c in sheet.get('cells',{}).items():
   row,col=position(ref)
   if row>nr or col>nc:raise ValueError('A cell is outside the sheet size')
   value=c.get('v','');typ=c.get('t','s');style=c.get('s',0)
   if not isinstance(value,str) or len(value)>4000 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',value) or typ not in ('s','n','b','f','d','inlineStr','e') or str(style) not in style_ids:raise ValueError('Invalid cell value or style')
   fmt={k:v for k,v in c.get('fmt',{}).items() if k in ('bold','italic','size','color','fill','align','wrap','border','number_format')}
   for k,v in fmt.items():
    if k in ('bold','italic','wrap','border') and type(v)!=bool:raise ValueError('Invalid cell format')
    if k=='size' and (not isinstance(v,(int,float)) or not 6<=v<=48):raise ValueError('Font size must be 6–48')
    if k in ('color','fill') and not re.fullmatch(r'#[0-9A-Fa-f]{6}',str(v)):raise ValueError('Invalid colour')
    if k=='align' and v not in ('left','center','right'):raise ValueError('Invalid alignment')
    if k=='number_format' and v not in ('General','0','0.00','0%','0.00%','yyyy-mm-dd','dd/mm/yyyy'):raise ValueError('Invalid number format')
   if typ=='n':
    try:
     if not math.isfinite(float(value)):raise ValueError()
    except ValueError:raise ValueError('Invalid numeric cell '+ref)
   if typ=='f' and (not value.startswith('=') or len(value)>1000):raise ValueError('Invalid formula at '+ref)
   if typ=='d':
    try:dt.datetime.fromisoformat(value)
    except ValueError:raise ValueError('Invalid date at '+ref)
   if typ=='b' and value not in ('True','False','1','0'):raise ValueError('Invalid boolean at '+ref)
   cells[ref]=dict(v=value,t=typ,s=int(style),**({'fmt':fmt} if fmt else {}))
  total+=len(cells)
  if total>50000:raise ValueError('Use up to 50,000 populated or formatted cells')
  covered=set();merges=sheet.get('merges',[])
  if len(merges)>1000:raise ValueError('Too many merged ranges')
  for area in merges:
   if not re.fullmatch(r'[A-Z]{1,3}[1-9]\d{0,3}:[A-Z]{1,3}[1-9]\d{0,3}',area):raise ValueError('Invalid merged range')
   a,b,x,y=range_boundaries(area)
   if not 1<=a<=x<=nc or not 1<=b<=y<=nr or (x-a+1)*(y-b+1)>10000:raise ValueError('Merged range is outside the sheet')
   for row in range(b,y+1):
    for col in range(a,x+1):
     ref=get_column_letter(col)+str(row)
     if ref in covered:raise ValueError('Merged ranges overlap')
     covered.add(ref)
     if (row,col)!=(b,a) and cells.get(ref,{}).get('v'):raise ValueError('Clear other values before merging '+area)
  sizes={}
  for key,maximum,lo,hi in [('widths',nc,35,600),('heights',nr,20,500)]:
   sizes[key]={}
   for k,v in sheet.get(key,{}).items():
    if not str(k).isdigit() or not 1<=int(k)<=maximum or not isinstance(v,(int,float)) or not lo<=v<=hi:raise ValueError('Invalid row / column size')
    sizes[key][str(k)]=v
  out.append(dict(id=sid,source=source,name=name,rows=nr,cols=nc,cells=cells,merges=merges,orientation='landscape' if sheet.get('orientation')=='landscape' else 'portrait',**sizes))
 return dict(sheets=out,styles=original['styles'])

def calculate(model):
 """Small bounded arithmetic interpreter; never eval workbook formulas or fetch links."""
 byname={s['name']:s for s in model['sheets']};memo={};active=set();budget=[0]
 ref_pattern=r"(?:(?:'((?:[^']|'')+)'|([A-Za-z_][A-Za-z0-9_]*))!)?(\$?[A-Z]{1,3}\$?[1-9]\d*)(?::(\$?[A-Z]{1,3}\$?[1-9]\d*))?(?![A-Za-z0-9_])"
 def cell(sheet,ref):
  key=(sheet,ref);budget[0]+=1
  if budget[0]>100000:raise ValueError('#LIMIT!')
  if key in memo:return memo[key]
  if key in active:raise ValueError('#CYCLE!')
  if len(active)>50:raise ValueError('#LIMIT!')
  if sheet not in byname:raise ValueError('#REF!')
  position(ref);c=byname[sheet]['cells'].get(ref,{});v=c.get('v','')
  if not v:return 0
  if c.get('t')=='n':return float(v)
  if c.get('t')!='f':return v
  active.add(key)
  try:
   def replace(m):return 'RANGE('+repr((m[1] or m[2] or sheet).replace("''","'"))+','+repr(m[3].replace('$',''))+','+repr((m[4] or m[3]).replace('$',''))+')' if m[4] else 'CELL('+repr((m[1] or m[2] or sheet).replace("''","'"))+','+repr(m[3].replace('$',''))+')'
   expr=re.sub(ref_pattern,replace,v[1:]);tree=ast.parse(expr.replace('^','**'),mode='eval')
   if sum(1 for _ in ast.walk(tree))>400:raise ValueError('#LIMIT!')
   def visit(n):
    if isinstance(n,ast.Constant) and isinstance(n.value,(int,float,str)):return n.value
    if isinstance(n,ast.UnaryOp) and isinstance(n.op,(ast.UAdd,ast.USub)):return visit(n.operand)*(1 if isinstance(n.op,ast.UAdd) else -1)
    if isinstance(n,ast.BinOp):
     a,b=visit(n.left),visit(n.right)
     if not all(isinstance(v,(int,float)) for v in (a,b)):raise ValueError('#VALUE!')
     if isinstance(n.op,ast.Add):return a+b
     if isinstance(n.op,ast.Sub):return a-b
     if isinstance(n.op,ast.Mult):return a*b
     if isinstance(n.op,ast.Div):return a/b
     if isinstance(n.op,ast.Pow) and abs(b)<=20 and abs(a)<=1e12:return a**b
    if isinstance(n,ast.Call) and isinstance(n.func,ast.Name) and not n.keywords:
     name=n.func.id.upper();args=[visit(a) for a in n.args]
     if name=='CELL' and len(args)==2:return cell(*args)
     if name=='RANGE' and len(args)==3:
      a,b=position(args[1]);x,y=position(args[2])
      if x<a or y<b or (x-a+1)*(y-b+1)>10000:raise ValueError('#LIMIT!')
      return [cell(args[0],get_column_letter(col)+str(row)) for row in range(a,x+1) for col in range(b,y+1)]
     vals=[v for arg in args for v in (arg if isinstance(arg,list) else [arg]) if isinstance(v,(int,float))]
     if name=='SUM':return sum(vals)
     if name=='COUNT':return len(vals)
     if name=='AVERAGE' and vals:return sum(vals)/len(vals)
     if name=='MIN' and vals:return min(vals)
     if name=='MAX' and vals:return max(vals)
     if name=='ABS' and len(vals)==1:return abs(vals[0])
     if name=='ROUND' and len(vals)==2 and -10<=vals[1]<=10:return round(vals[0],int(vals[1]))
    raise ValueError('#EXCEL!')
   value=visit(tree.body)
   if not isinstance(value,(int,float)) or not math.isfinite(value):raise ValueError('#VALUE!')
   memo[key]=value;return value
  finally:active.remove(key)
 result={}
 for sheet in model['sheets']:
  result[sheet['id']]={}
  for ref,c in sheet['cells'].items():
   if c.get('t')=='f':
    try:result[sheet['id']][ref]=cell(sheet['name'],ref)
    except Exception as e:result[sheet['id']][ref]=str(e) if str(e).startswith('#') else '#EXCEL!'
 return result

def export(raw,model):
 from openpyxl.styles.stylesheet import write_stylesheet
 from openpyxl.xml.functions import tostring
 data=parts(raw)
 with warnings.catch_warnings():
  warnings.simplefilter('ignore',UserWarning);original=load_workbook(io.BytesIO(raw),keep_links=False)
 original_sizes={}
 try:
  # Loading merged cells generates additional border styles. Include those IDs,
  # without resaving the workbook and losing its unsupported vector drawings.
  for ws in original:
   original_sizes[ws.title]=sizes(ws)
   for row in ws:
    for cell in row:_=cell.style_id
  styles=xml(tostring(write_stylesheet(original)))
 finally:original.close()
 fonts=styles.find(tag('fonts'));fills=styles.find(tag('fills'));borders=styles.find(tag('borders'));xfs=styles.find(tag('cellXfs'));cache={}
 def append(parent,node):parent.append(node);parent.set('count',str(len(parent)));return len(parent)-1
 def style(c):
  base=c.get('s',0);fmt=c.get('fmt',{})
  if not fmt:return base
  key=(base,json.dumps(fmt,sort_keys=True))
  if key in cache:return cache[key]
  xf=copy.deepcopy(xfs[base]);font=copy.deepcopy(fonts[int(xf.get('fontId','0'))])
  for name,k in [('b','bold'),('i','italic'),('sz','size'),('color','color')]:
   if k not in fmt:continue
   old=font.find(tag(name))
   if old is not None:font.remove(old)
   if k in ('bold','italic'):
    if fmt[k]:ET.SubElement(font,tag(name))
   else:ET.SubElement(font,tag(name),{'rgb':'FF'+fmt[k][1:]} if k=='color' else {'val':str(fmt[k])})
  xf.set('fontId',str(append(fonts,font)));xf.set('applyFont','1')
  if 'fill' in fmt:
   fill=ET.Element(tag('fill'));pattern=ET.SubElement(fill,tag('patternFill'),{'patternType':'solid'});ET.SubElement(pattern,tag('fgColor'),{'rgb':'FF'+fmt['fill'][1:]});ET.SubElement(pattern,tag('bgColor'),{'indexed':'64'});xf.set('fillId',str(append(fills,fill)));xf.set('applyFill','1')
  if 'border' in fmt:
   border=ET.Element(tag('border'))
   for edge in ('left','right','top','bottom'):ET.SubElement(border,tag(edge),{'style':'thin'} if fmt['border'] else {})
   xf.set('borderId',str(append(borders,border)));xf.set('applyBorder','1')
  alignment=xf.find(tag('alignment'))
  if alignment is None:alignment=ET.SubElement(xf,tag('alignment'))
  if 'align' in fmt:alignment.set('horizontal',fmt['align'])
  if 'wrap' in fmt:alignment.set('wrapText','1' if fmt['wrap'] else '0')
  if 'number_format' in fmt:
   from openpyxl.styles.numbers import BUILTIN_FORMATS_REVERSE
   code=fmt['number_format'];num_id=BUILTIN_FORMATS_REVERSE.get(code)
   if num_id is None:
    formats=styles.find(tag('numFmts'))
    if formats is None:formats=ET.Element(tag('numFmts'),{'count':'0'});styles.insert(0,formats)
    num_id=next((int(n.get('numFmtId')) for n in formats if n.get('formatCode')==code),None)
    if num_id is None:
     num_id=max([163]+[int(n.get('numFmtId')) for n in formats])+1;ET.SubElement(formats,tag('numFmt'),{'numFmtId':str(num_id),'formatCode':code});formats.set('count',str(len(formats)))
   xf.set('numFmtId',str(num_id));xf.set('applyNumberFormat','1')
  xf.set('applyAlignment','1');cache[key]=append(xfs,xf);return cache[key]
 wb=xml(data['xl/workbook.xml']);wb_sheets=wb.find(tag('sheets'));rels=xml(data['xl/_rels/workbook.xml.rels']);types=xml(data['[Content_Types].xml']);paths=relationships(data,'xl/workbook.xml');nodes={paths[s.get('{'+R+'}id')]:s for s in wb_sheets};computed=calculate(model)
 # Rename print-area/sheet references in defined names, preserving the original package.
 renamed={n.get('name'):s['name'] for s in model['sheets'] if (n:=nodes.get(s['source'])) is not None and n.get('name')!=s['name']}
 for node in wb.findall(tag('definedNames')+'/'+tag('definedName')):
  if node.text:
   for old,new in renamed.items():node.text=node.text.replace("'"+old.replace("'","''")+"'!","'"+new.replace("'","''")+"'!")
 for index,sheet in enumerate(model['sheets']):
  source=sheet['source'];path=source or 'xl/worksheets/workshop-'+sheet['id']+'.xml'
  if source:root=xml(data[source]);nodes[source].set('name',sheet['name'])
  else:
   rid='workshop-'+sheet['id'];root=ET.Element(tag('worksheet'));ET.SubElement(root,tag('sheetData'));ET.SubElement(wb_sheets,tag('sheet'),{'name':sheet['name'],'sheetId':str(max(int(n.get('sheetId')) for n in wb_sheets)+1),'{'+R+'}id':rid});ET.SubElement(rels,'{'+P+'}Relationship',{'Id':rid,'Type':R+'/worksheet','Target':'worksheets/workshop-'+sheet['id']+'.xml'});ET.SubElement(types,'{'+C+'}Override',{'PartName':'/'+path,'ContentType':'application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml'})
  sheetdata=root.find(tag('sheetData'));row_nodes={int(r.get('r')):r for r in sheetdata}
  for row in row_nodes.values():
   for c in list(row):
    if c.tag==tag('c'):row.remove(c)
  for ref,c in sorted(sheet['cells'].items(),key=lambda item:position(item[0])):
   row,col=position(ref)
   if row not in row_nodes:row_nodes[row]=ET.SubElement(sheetdata,tag('row'),{'r':str(row)})
   cell=ET.SubElement(row_nodes[row],tag('c'),{'r':ref,'s':str(style(c))});v=c['v'];typ=c['t']
   if typ=='f':
    ET.SubElement(cell,tag('f')).text=v[1:];result=computed[sheet['id']].get(ref)
    if isinstance(result,(int,float)):ET.SubElement(cell,tag('v')).text=str(result)
   elif typ in ('n','d','b') and v:
    cell.set('t',typ);ET.SubElement(cell,tag('v')).text=('1' if v in ('True','1') else '0') if typ=='b' else v
   else:
    cell.set('t','inlineStr');textnode=ET.SubElement(ET.SubElement(cell,tag('is')),tag('t'),{'{http://www.w3.org/XML/1998/namespace}space':'preserve'});textnode.text=v
  original_name=next((old for old,new in renamed.items() if new==sheet['name']),sheet['name']);old_widths,old_heights=original_sizes.get(original_name,({},{})) if source else ({},{})
  for number,height in sheet['heights'].items():
   if old_heights.get(number)==height:continue
   number=int(number)
   if number not in row_nodes:row_nodes[number]=ET.SubElement(sheetdata,tag('row'),{'r':str(number)})
   row_nodes[number].set('ht',str(height*3/4));row_nodes[number].set('customHeight','1')
  sheetdata[:]=sorted(sheetdata,key=lambda n:int(n.get('r')))
  cols=root.find(tag('cols'))
  if cols is None:cols=ET.Element(tag('cols'));root.insert(list(root).index(sheetdata),cols)
  for col,width in sheet['widths'].items():
   if old_widths.get(col)!=width:ET.SubElement(cols,tag('col'),{'min':col,'max':col,'width':str(width/7),'customWidth':'1'})
  # Avoid overlapping column definitions; latest per-column width wins while preserving hidden/style metadata.
  column_nodes={}
  for col in cols:
   for n in range(int(col.get('min')),min(int(col.get('max')),sheet['cols'])+1):
    attrs={**(column_nodes.get(n,{})),**col.attrib,'min':str(n),'max':str(n)};column_nodes[n]=attrs
  cols[:]=[ET.Element(tag('col'),attrs) for n,attrs in sorted(column_nodes.items())]
  merged=root.find(tag('mergeCells'))
  if merged is None:merged=ET.Element(tag('mergeCells'));root.insert(list(root).index(sheetdata)+1,merged)
  merged[:]=[ET.Element(tag('mergeCell'),{'ref':area}) for area in sheet['merges']];merged.set('count',str(len(merged)))
  setup=root.find(tag('pageSetup'))
  if setup is not None:setup.set('orientation',sheet['orientation'])
  dim=root.find(tag('dimension'))
  if dim is not None:dim.set('ref','A1:'+get_column_letter(sheet['cols'])+str(sheet['rows']))
  data[path]=encoded(root)
 calc=wb.find(tag('calcPr'))
 if calc is None:calc=ET.SubElement(wb,tag('calcPr'))
 calc.set('fullCalcOnLoad','1');calc.set('forceFullCalc','1')
 data['xl/workbook.xml']=encoded(wb);data['xl/_rels/workbook.xml.rels']=encoded(rels);data['[Content_Types].xml']=encoded(types);data['xl/styles.xml']=encoded(styles)
 out=io.BytesIO()
 with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
  for name,value in data.items():z.writestr(name,value)
 return out.getvalue()

def view_images(raw,model):
 data=parts(raw);out={};notes=[]
 for s in model['sheets']:
  images,issues=pictures(data,s['source']) if s['source'] else ([],[]);out[s['id']]=images;notes.extend(issues)
 return out,list(dict.fromkeys(notes))
