"""QC-owned dispatch batches, traceability and reviewed measurement reports."""
import copy,datetime as dt,hashlib,io,json,zipfile
import access,manufacturing as m,hr_ops,qc_excel as x
from storage import identifier

COLLECTIONS=('qc_batches','qc_reports','qc_trace_lots','qc_imports','qc_audit')
def today():return dt.datetime.now(dt.timezone(dt.timedelta(hours=5,minutes=30))).date().isoformat()
def project(c,jid,user,edit=False,author=False):
 j=m.get_job(c,jid,user)
 if edit:
  if j.get('flow_version')!=2:raise ValueError('Use a plant project with the QC workflow for dispatch reports')
  access.require(user,access.ROLES if author else ('QC',))
  if not author:hr_ops.assigned(c,j,user)
  if j['stage']=='Completed':raise ValueError('Completed projects are read-only in QC')
 return j
def batch(c,d,user,edit=False,author=False):
 b=c.one('qc_batches',{'id':identifier(d.get('batch_id'))})
 if not b:raise ValueError('QC batch not found')
 j=project(c,b['job_id'],user,edit,author)
 if edit:
  if b['version']!=d.get('version'):raise ValueError('QC batch changed. Reload it before saving.')
  if b.get('actual_date'):raise ValueError('Dispatched batches are read-only; create a separate batch for further work')
 return b,j
def log(c,b,user,action,notes='',changes=None):
 c.insert('qc_audit',job_id=b['job_id'],batch_id=b.get('id'),action=action,notes=str(notes)[:4000],actor=user['name'],actor_id=user['id'],created_at=m.now(),changes=changes or [])
 m.history(c,b['job_id'],user,action,notes);m.bump(c,b['job_id'])
 if b.get('id'):c.update('qc_batches',{'id':b['id']},inc={'version':1})
def invalidate(c,b,j):
 steps=copy.deepcopy(b['steps']);steps['Final QC']=dict(status='Pending',quantity=0,notes='QC evidence changed');steps['Packing']=dict(status='Pending',quantity=0,notes='Final QC required')
 c.update('qc_batches',{'id':b['id']},{'steps':steps},inc={'evidence_version':1})
 c.update('work_jobs',{'id':j['id']},{'final_inspection':False,**({'stage':'QC inspection'} if j['stage']=='Packing & dispatch' else {})})
def blank_steps(source=None):return {s:dict(status='NA' if (source or {}).get(s,'').upper()=='NA' and s in ('Machining','Tapping') else 'Pending',quantity=0,notes='') for s in x.STEPS}
def create_batch(c,d,user):
 j=project(c,d.get('job_id'),user,True,author=True)
 quantity=x.integer(d.get('quantity'),100)
 values=dict(job_id=j['id'],plant_id=j.get('plant_id',''),po=j['po_number'],customer=j['customer'],drawing=m.text(d,'drawing')[:100],description=str(d.get('description',''))[:200],thickness=m.text(d,'thickness')[:80],quantity=quantity,tc_required=d.get('tc_required') is True,planned_date=dt.date.fromisoformat(str(d.get('planned_date'))).isoformat(),actual_date='',source_actual_date='',drawing_ref='',numbering_ref='',source_steps={},steps=blank_steps(),allocations=[],document_ids=[],current_report=None,version=1,created_at=m.now())
 bid=c.insert('qc_batches',**values)
 c.update('work_jobs',{'id':j['id']},{'final_inspection':False,**({'stage':'QC inspection'} if j['stage']=='Packing & dispatch' else {})})
 log(c,{**values,'id':bid},user,'QC dispatch batch created',values['drawing']);return bid
def schedule(c,d,user):
 b,j=batch(c,d,user,True,author=True)
 c.update('qc_batches',{'id':b['id']},{'planned_date':dt.date.fromisoformat(str(d.get('planned_date'))).isoformat(),'description':str(d.get('description',b['description']))[:200]})
 log(c,b,user,'QC dispatch schedule updated',m.text(d,'notes'))
def preview(c,d,user):
 access.require(user,('QC',) if d.get('kind')=='tracker' else access.ROLES)
 return x.tracker(d.get('file')) if d.get('kind')=='tracker' else x.inspection(d.get('file'),blank=d.get('mode')!='historical')
def import_tracker(c,d,user):
 j=project(c,d.get('job_id'),user,True);parsed=x.tracker(d.get('file'));raw=x.file_bytes(d['file']);digest=hashlib.sha256(raw).hexdigest()
 if c.one('qc_imports',{'job_id':j['id'],'sha256':digest}):raise ValueError('This workbook is already imported into this project')
 if any(b['po']!=str(j['po_number']).strip() for b in parsed['batches']):raise ValueError('Workbook PO does not match the selected project. Choose the correct PO project.')
 did=m.upload(c,j['id'],user,'Supporting',d['file'])
 iid=c.insert('qc_imports',job_id=j['id'],sha256=digest,document_id=did,warnings=parsed['warnings'],imported_by=user['name'],imported_at=m.now())
 for source in parsed['batches']:
  c.insert('qc_batches',**source,job_id=j['id'],plant_id=j.get('plant_id',''),customer=j['customer'],tc_required=False,import_id=iid,actual_date='',steps=blank_steps(source['source_steps']),allocations=[],document_ids=[],current_report=None,version=1,created_at=m.now())
 for trace in parsed['trace_lots']:c.insert('qc_trace_lots',**trace,job_id=j['id'],import_id=iid)
 c.update('work_jobs',{'id':j['id']},{'final_inspection':False,**({'stage':'QC inspection'} if j['stage']=='Packing & dispatch' else {})})
 log(c,{'job_id':j['id']},user,'QC tracker imported as unverified batches',str(len(parsed['batches']))+' drawing / dispatch rows; '+str(len(parsed['trace_lots']))+' unallocated material rows')
def save_steps(c,d,user):
 b,j=batch(c,d,user,True);steps=d.get('steps')
 if not isinstance(steps,dict) or set(steps)!=set(x.STEPS):raise ValueError('All process stages are required')
 cleaned={}
 for s in x.STEPS:
  v=steps[s];status=v.get('status');qty=v.get('quantity');notes=str(v.get('notes','')).strip()[:1000]
  if status not in ('Pending','In progress','Done','NA','Rework'):raise ValueError('Invalid process status')
  if isinstance(qty,bool) or not isinstance(qty,int) or not 0<=qty<=b['quantity']:raise ValueError('Stage quantity must be a whole number within batch quantity')
  if status=='Done' and qty!=b['quantity']:raise ValueError('Use In progress for partial quantities: '+s)
  if status in ('Pending','NA') and qty:raise ValueError('Pending / NA stages must have zero completed quantity')
  if status=='NA' and s in ('QC','Final QC','Packing'):raise ValueError(s+' cannot be skipped')
  if status in ('NA','Rework') and not notes:raise ValueError('Enter a reason for '+s+' '+status)
  if status=='Done' and any(v['status'] not in ('Done','NA') for v in cleaned.values()):raise ValueError('Complete earlier stages before '+s)
  cleaned[s]=dict(status=status,quantity=qty,notes=notes)
 if cleaned['Final QC']['status']=='Done':
  errors=release_errors(c,{**b,'steps':cleaned},check_steps=False)
  if errors:raise ValueError('; '.join(errors))
 if cleaned['Packing']['status']=='Done' and not j.get('final_inspection'):raise ValueError('Complete project final inspection before confirming packing')
 if cleaned!=b['steps']:
  c.update('qc_batches',{'id':b['id']},{'steps':cleaned})
  if cleaned['Final QC']['status']!='Done':c.update('work_jobs',{'id':j['id']},{'final_inspection':False,**({'stage':'QC inspection'} if j['stage']=='Packing & dispatch' else {})})
  log(c,b,user,'QC process progress updated',m.text(d,'notes'))
def trace(c,d,user):
 b,j=batch(c,d,user,True)
 if d.get('new_lot'):
  v=d['new_lot'];c.insert('qc_trace_lots',job_id=j['id'],heat_no=m.text(v,'heat_no')[:100],plate_no=m.text(v,'plate_no')[:100],ai_no=str(v.get('ai_no',''))[:100],thickness=b['thickness'],quantity=x.integer(v.get('quantity')),planned_date=b['planned_date'],source_actual_date='');log(c,b,user,'Material trace lot created');return
 rows=d.get('allocations')
 if not isinstance(rows,list) or len(rows)>100:raise ValueError('Choose material lots and quantities')
 allocations=[];seen=set()
 for r in rows:
  lid=identifier(r.get('lot_id'));qty=x.integer(r.get('quantity'));lot=c.one('qc_trace_lots',{'id':lid,'job_id':j['id']})
  if not lot or lid in seen:raise ValueError('Choose unique material lots from this project')
  seen.add(lid)
  if lot['thickness'].replace(' ','').lower()!=b['thickness'].replace(' ','').lower():raise ValueError('Material thickness does not match the drawing batch')
  used=sum(a['quantity'] for other in c.all('qc_batches',{'job_id':j['id'],'id':{'$ne':b['id']}}) for a in other.get('allocations',[]) if a['lot_id']==lid)
  if used+qty>lot['quantity']:raise ValueError('Material allocation exceeds available quantity for heat '+lot.get('heat_no',''))
  allocations.append(dict(lot_id=lid,quantity=qty))
 if sum(r['quantity'] for r in allocations)>b['quantity']:raise ValueError('Traceability allocation exceeds batch quantity')
 c.update('qc_batches',{'id':b['id']},{'allocations':allocations});invalidate(c,b,j);log(c,b,user,'Heat / plate traceability updated')
def document(c,d,user):
 b,j=batch(c,d,user,True);kind=d.get('kind')
 if kind not in ('Drawing','Numbering','TC','EWAY challan'):raise ValueError('Choose Drawing, Numbering, TC or EWAY challan')
 if d.get('document_id'):
  doc=c.one('job_documents',{'id':identifier(d['document_id']),'job_id':j['id']})
  if not doc or doc.get('qc_kind',doc['kind'])!=kind:raise ValueError('Choose a matching document from this project')
  did=doc['id']
 else:
  did=m.upload(c,j['id'],user,'Supporting' if kind=='Numbering' else kind,d.get('file'));c.update('job_documents',{'id':did},{'qc_kind':kind})
 ids=list(dict.fromkeys(b['document_ids']+[did]));c.update('qc_batches',{'id':b['id']},{'document_ids':ids})
 if kind!='EWAY challan':invalidate(c,b,j)
 log(c,b,user,'QC batch document linked',kind)
def import_report(c,d,user):
 b,j=batch(c,d,user,True,author=True)
 mode=d.get('mode','blank')
 if mode not in ('blank','historical'):raise ValueError('Choose blank template or historical draft')
 if d.get('template_report_id'):
  source=c.one('qc_reports',{'id':identifier(d['template_report_id'])})
  if not source:raise ValueError('Inspection format not found')
  project(c,source['job_id'],user);parsed=copy.deepcopy(source);mode='blank';did=None
 else:
  parsed=x.inspection(d.get('file'),blank=mode=='blank')['report'];did=m.upload(c,j['id'],user,'Supporting',d['file'])
 if parsed['drawing']!=b['drawing']:raise ValueError('Inspection drawing number does not match this batch')
 if mode=='historical' and (parsed['po']!=b['po'] or parsed['quantity']!=b['quantity']):raise ValueError('Historical report PO and quantity must match the batch')
 chars=copy.deepcopy(parsed['characteristics'])
 if mode=='blank':
  for r in chars:r.update(values=['']*b['quantity'],manual_result='',manual_reason='')
 rid=c.insert('qc_reports',job_id=j['id'],batch_id=b['id'],drawing=b['drawing'],po=b['po'],customer=b['customer'],part_name=parsed['part_name'],revision=parsed['revision'],quantity=b['quantity'],inspection_date=today() if mode=='blank' else parsed['inspection_date'],status='Draft',version=1,characteristics=chars,inspector=user['name'],source_document_id=did,source_mode=mode,warnings=[] if mode=='blank' else parsed.get('warnings',[]),created_at=m.now(),review_notes='')
 c.update('qc_reports',{'id':rid},{'reference':f'QC-{rid:05}', 'summary':x.results({'characteristics':chars})})
 c.update('qc_batches',{'id':b['id']},{'current_report':rid});invalidate(c,b,j);log(c,b,user,'Inspection draft created',f'QC-{rid:05} · '+mode)
def save_report(c,d,user):
 b,j=batch(c,d,user,True,author=True);r=c.one('qc_reports',{'id':b['current_report'],'batch_id':b['id']})
 if not r or r['status']!='Draft':raise ValueError('Create a new draft revision before changing an approved report')
 if d.get('report_version')!=r['version']:raise ValueError('Inspection readings changed. Reload the report.')
 incoming=d.get('characteristics')
 if not isinstance(incoming,list) or len(incoming)!=len(r['characteristics']):raise ValueError('Inspection characteristic count changed')
 chars=[]
 for original,values in zip(r['characteristics'],incoming):
  readings=values.get('values')
  if not isinstance(readings,list) or len(readings)!=r['quantity'] or any(not isinstance(v,str) or len(v)>100 for v in readings):raise ValueError('Enter one reading per piece (maximum 100 characters)')
  manual=values.get('manual_result','');reason=str(values.get('manual_reason','')).strip()[:1000]
  if manual not in ('','OK','Rework') or (manual and (original.get('limits') or not reason)):raise ValueError('Manual assessment requires an unrecognised specification and a QC reason')
  chars.append({**original,'values':readings,'manual_result':manual,'manual_reason':reason,'manual_by':user['name'] if manual else ''})
 values=dict(characteristics=chars,summary=x.results({'characteristics':chars}),inspection_date=dt.date.fromisoformat(str(d.get('inspection_date'))).isoformat(),review_notes=str(d.get('review_notes',''))[:4000],inspector=user['name'])
 if values['inspection_date']>today():raise ValueError('Inspection date cannot be in the future')
 changes=[]
 for old,new in zip(r['characteristics'],chars):
  for i,(before,after) in enumerate(zip(old['values'],new['values'])):
   if before!=after:changes.append(dict(characteristic=old['number'],piece=i+1,before=before,after=after))
  if old.get('manual_result')!=new['manual_result'] or old.get('manual_reason')!=new['manual_reason']:changes.append(dict(characteristic=old['number'],manual_before=[old.get('manual_result',''),old.get('manual_reason','')],manual_after=[new['manual_result'],new['manual_reason']]))
 c.update('qc_reports',{'id':r['id']},values,inc={'version':1});invalidate(c,b,j);log(c,b,user,'Inspection readings saved',r['reference'],changes=changes)
def approve_report(c,d,user):
 b,j=batch(c,d,user,True);r=c.one('qc_reports',{'id':b['current_report']})
 if not r or r['status']!='Draft':raise ValueError('Choose a draft inspection report')
 if d.get('report_version')!=r['version']:raise ValueError('Report changed. Reload it before approval.')
 summary=x.results(r)
 if summary['Missing'] or summary['Rework'] or summary['Review']:raise ValueError('Resolve missing, out-of-tolerance and unreviewed measurements before approval')
 if not r.get('inspection_date') or r['inspection_date']>today():raise ValueError('Enter a valid inspection date before approval')
 evidence=release_errors(c,b,check_steps=False)
 evidence=[error for error in evidence if error!='Approve the current inspection report']
 if evidence:raise ValueError('; '.join(evidence))
 c.update('qc_reports',{'id':r['id']},{'status':'Approved','summary':summary,'approved_by':user['name'],'approved_user_id':user['id'],'approved_at':m.now(),'approval_notes':m.text(d,'notes'),'snapshot_batch':b,'snapshot_trace':traces(c,b),'snapshot_company':(c.one('settings',{'key':'company'}) or {}).get('value','TPPL')},inc={'version':1})
 log(c,b,user,'Inspection report approved',r['reference']+' · '+d['notes']);return
def revise_report(c,d,user):
 b,j=batch(c,d,user,True,author=True);r=c.one('qc_reports',{'id':b.get('current_report')})
 if not r:raise ValueError('Import an inspection report first')
 if r['status']=='Draft':raise ValueError('Save the existing draft before creating a revision')
 fields=('drawing','po','customer','part_name','revision','quantity','characteristics','warnings')
 values={k:copy.deepcopy(r.get(k,[] if k in ('characteristics','warnings') else '')) for k in fields}
 rid=c.insert('qc_reports',**values,job_id=j['id'],batch_id=b['id'],status='Draft',version=1,inspection_date=today(),inspector=user['name'],created_at=m.now(),review_notes=m.text(d,'notes'),revises=r['id'],summary=x.results(r))
 c.update('qc_reports',{'id':rid},{'reference':f'QC-{rid:05}'})
 c.update('qc_batches',{'id':b['id']},{'current_report':rid});invalidate(c,b,j);log(c,b,user,'Inspection revision created',r['reference']+' → '+f'QC-{rid:05}')
def traces(c,b):
 out=[]
 for a in b.get('allocations',[]):
  lot=c.one('qc_trace_lots',{'id':a['lot_id'],'job_id':b['job_id']})
  if lot:out.append({**lot,'allocated':a['quantity']})
 return out
def release_errors(c,b,check_steps=True,include_packing=True):
 errors=[];report=c.one('qc_reports',{'id':b.get('current_report')}) if b.get('current_report') else None
 if not report or report['status']!='Approved':errors.append('Approve the current inspection report')
 elif report.get('snapshot_batch',{}).get('evidence_version',0)!=b.get('evidence_version',0):errors.append('Evidence changed: create and approve an inspection revision')
 if check_steps and any(b['steps'][s]['status'] not in ('Done','NA') for s in x.STEPS if include_packing or s!='Packing'):errors.append('Complete all process / packing stages' if include_packing else 'Complete process stages through Final QC')
 lots=traces(c,b)
 if sum(r['allocated'] for r in lots)!=b['quantity'] or any(not r.get('heat_no') or not r.get('plate_no') for r in lots):errors.append('Allocate complete heat / plate traceability for every piece')
 docs=c.all('job_documents',{'id':{'$in':b['document_ids']},'job_id':b['job_id']});kinds={d.get('qc_kind',d['kind']) for d in docs}
 for kind in (('Drawing','TC') if b.get('tc_required') else ('Drawing',)):
  if kind not in kinds:errors.append('Attach '+kind)
 if b.get('numbering_ref') and 'Numbering' not in kinds:errors.append('Attach the numbering reference')
 if not b.get('planned_date'):errors.append('Confirm the planned dispatch date')
 return errors
def dispatch(c,d,user):
 b,j=batch(c,d,user,True)
 if j['stage']!='Packing & dispatch' or not j.get('final_inspection'):raise ValueError('Complete project final inspection before recording batch dispatch')
 if c.one('project_requests',{'job_id':j['id'],'status':{'$in':['Open','Pending']}}):raise ValueError('Resolve department requests and QC approvals before dispatch')
 if c.one('inspection_calls',{'job_id':j['id'],'status':{'$in':['Pending Admin','Approved','Rework']}}):raise ValueError('Resolve the customer inspection call before dispatch')
 errors=release_errors(c,b)
 if errors:raise ValueError('; '.join(errors))
 docs=c.all('job_documents',{'id':{'$in':b['document_ids']},'job_id':j['id']})
 if not any(doc.get('qc_kind',doc['kind'])=='EWAY challan' for doc in docs):raise ValueError('Attach the dispatch EWAY challan')
 actual=dt.date.fromisoformat(str(d.get('actual_date'))).isoformat()
 if actual>today():raise ValueError('Actual dispatch date cannot be in the future')
 c.update('qc_batches',{'id':b['id']},{'actual_date':actual,'dispatch_notes':m.text(d,'notes'),'dispatched_by':user['name']});log(c,b,user,'QC batch dispatched',actual)
def detail(c,bid,user):
 b,j=batch(c,{'batch_id':bid},user)
 reports=c.all('qc_reports',{'batch_id':b['id']},sort=[('id',-1)],projection={'id':1,'reference':1,'revision':1,'status':1,'inspector':1,'approved_by':1})
 if b.get('current_report'):
  current=c.one('qc_reports',{'id':b['current_report'],'batch_id':b['id']})
  reports=[current if r['id']==b['current_report'] else r for r in reports]
 lots=c.all('qc_trace_lots',{'job_id':j['id']})
 others=c.all('qc_batches',{'job_id':j['id'],'id':{'$ne':b['id']}})
 for lot in lots:lot['available']=lot['quantity']-sum(a['quantity'] for other in others for a in other.get('allocations',[]) if a['lot_id']==lot['id'])
 return dict(batch=b,project=j,reports=reports,trace_lots=lots,documents=c.all('job_documents',{'job_id':j['id']},projection={'content':0}),history=c.all('qc_audit',{'batch_id':b['id']},sort=[('id',-1)],projection={'changes':0}),release_errors=release_errors(c,b))
def state(c,data,user):
 ids=[j['id'] for j in data['work_jobs']]
 data['qc_batches']=c.all('qc_batches',{'job_id':{'$in':ids}},sort=[('id',-1)])
 reports={r['id']:r for r in c.all('qc_reports',{'job_id':{'$in':ids}},projection={'id':1,'status':1,'summary':1})}
 for b in data['qc_batches']:
  r=reports.get(b.get('current_report'),{});b['inspection_status']=r.get('status','Not started');b['inspection_summary']=r.get('summary',{})
 data['qc_formats']=c.all('qc_reports',{'job_id':{'$in':ids}},projection={'id':1,'drawing':1,'revision':1,'reference':1,'part_name':1,'job_id':1},sort=[('id',-1)])
def gate(c,j,dispatching=False):
 for b in c.all('qc_batches',{'job_id':j['id']}):
  errors=release_errors(c,b,include_packing=dispatching)
  if dispatching and not b.get('actual_date'):errors.append('Record batch dispatch')
  if errors:raise ValueError('QC batch '+b['drawing']+' #'+str(b['id'])+': '+'; '.join(errors))
def export(c,rid,user,kind):
 r=c.one('qc_reports',{'id':identifier(rid)})
 if not r:raise ValueError('Inspection report not found')
 b,j=batch(c,{'batch_id':r['batch_id']},user);company=(c.one('settings',{'key':'company'}) or {}).get('value','TPPL')
 if kind not in ('xlsx','pdf'):raise ValueError('Choose PDF or Excel')
 content=(x.export_pdf if kind=='pdf' else x.export_xlsx)(r,r.get('snapshot_batch',b),r.get('snapshot_trace',traces(c,b)),r.get('snapshot_company',company))
 return {'_download':dict(name=r['reference']+'-'+r['drawing']+'.'+kind,content=content)}
def packet(c,bid,user):
 b,j=batch(c,{'batch_id':bid},user);buffer=io.BytesIO()
 reports=c.all('qc_reports',{'batch_id':b['id']});document_ids=set(b['document_ids'])
 document_ids.update(r['source_document_id'] for r in reports if r.get('source_document_id'))
 source=c.one('qc_imports',{'id':b['import_id'],'job_id':j['id']}) if b.get('import_id') else None
 if source:document_ids.add(source['document_id'])
 with zipfile.ZipFile(buffer,'w',zipfile.ZIP_DEFLATED) as z:
  z.writestr('batch-and-process.json',json.dumps(b,indent=2,default=str));z.writestr('traceability.json',json.dumps(traces(c,b),indent=2,default=str));z.writestr('history.json',json.dumps(c.all('qc_audit',{'batch_id':b['id']}),indent=2,default=str))
  for r in reports:
   z.writestr(r['reference']+'.json',json.dumps(r,indent=2,default=str))
   for kind in ('pdf','xlsx'):
    doc=export(c,r['id'],user,kind)['_download'];z.writestr(doc['name'],doc['content'])
  for doc in c.all('job_documents',{'id':{'$in':list(document_ids)},'job_id':j['id']}):z.writestr('documents/'+str(doc['id'])+'-'+doc['name'],doc['content'])
 raw=buffer.getvalue()
 if len(raw)>4*1024*1024:raise ValueError('Packet exceeds 4 MB. Download reports and documents individually, or use the Admin project archive.')
 return {'_download':dict(name='QC-batch-'+str(b['id'])+'.zip',content=raw)}

def tracker_export(c,jid,user):
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill
 from openpyxl.utils import get_column_letter
 j=project(c,jid,user);batches=c.all('qc_batches',{'job_id':j['id']},sort=[('planned_date',1),('id',1)]);wb=Workbook();ws=wb.active;ws.title='PROCESS'
 ws.append(['PO','Drawing No','Thickness',*x.STEPS,'Qty','Dispatch Date','Actual Dispatch Date','Drawing','Numbering Ref','Inspection status','Report','Batch ID'])
 for b in batches:
  r=c.one('qc_reports',{'id':b.get('current_report')}) if b.get('current_report') else None
  ws.append([b['po'],b['drawing'],b['thickness'],*[b['steps'][s]['status']+' ('+str(b['steps'][s]['quantity'])+')' for s in x.STEPS],b['quantity'],b['planned_date'],b['actual_date'],b['drawing_ref'],b['numbering_ref'],r['status'] if r else 'Not started',r['reference'] if r else '',b['id']])
 ws=wb.create_sheet('Allocated traceability');ws.append(['Batch ID','Drawing','Dispatch date','Thickness','Heat No','Plate No','AI No','Allocated quantity'])
 for b in batches:
  for t in traces(c,b):ws.append([b['id'],b['drawing'],b['planned_date'],t['thickness'],t['heat_no'],t['plate_no'],t['ai_no'],t['allocated']])
 ws=wb.create_sheet('Material lots');ws.append(['Lot ID','Thickness','Heat No','Plate No','AI No','Total quantity','Allocated','Unallocated'])
 for t in c.all('qc_trace_lots',{'job_id':j['id']}):
  used=sum(a['quantity'] for b in batches for a in b['allocations'] if a['lot_id']==t['id']);ws.append([t['id'],t['thickness'],t['heat_no'],t['plate_no'],t['ai_no'],t['quantity'],used,t['quantity']-used])
 for ws in wb:
  ws.freeze_panes='D2';ws.auto_filter.ref=ws.dimensions
  for row in ws:
   for cell in row:
    if isinstance(cell.value,str):cell.data_type='s'
  for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17664F')
  for i in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(i)].width=24
 output=io.BytesIO();wb.save(output)
 return {'_download':dict(name='PO-'+str(j['id'])+'-QC-dispatch-tracker.xlsx',content=output.getvalue())}

HANDLERS={'create':create_batch,'schedule':schedule,'import-tracker':import_tracker,'steps':save_steps,'trace':trace,'document':document,'import-report':import_report,'save-report':save_report,'approve-report':approve_report,'revise-report':revise_report,'dispatch':dispatch}
